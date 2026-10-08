from __future__ import annotations

from array import array
from dataclasses import dataclass, field
from datetime import datetime, timezone
from io import BytesIO
from functools import lru_cache
from threading import Lock
import hashlib
import math
import os
import re
from typing import Any, Callable, Iterable

from . import db

EXTRACTOR_VERSION = "local-extractors/1"
CHUNKER_VERSION = "paragraph-window/1"
DEFAULT_EMBEDDING_MODEL = os.environ.get(
    "KNOWLEDGE_EMBEDDING_MODEL", "sentence-transformers/all-MiniLM-L6-v2"
)
DEFAULT_CAPTION_MODEL = os.environ.get(
    "KNOWLEDGE_CAPTION_MODEL", "Salesforce/blip-image-captioning-base"
)
CHUNK_CHARS = 1_200
CHUNK_OVERLAP_CHARS = 180
TOKEN_PATTERN = re.compile(r"[\w'-]+", re.UNICODE)
REBUILD_LOCK = Lock()


class OptionalDependencyError(RuntimeError):
    pass


@dataclass
class Segment:
    text: str
    locator: dict[str, Any] = field(default_factory=dict)


@dataclass
class Extraction:
    segments: list[Segment]
    metadata: dict[str, Any]
    extractor_id: str

    @property
    def text(self) -> str:
        return "\n\n".join(segment.text for segment in self.segments if segment.text.strip())


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> Extraction:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise OptionalDependencyError("PDF extraction requires pypdf") from exc
    reader = PdfReader(BytesIO(data))
    segments = [Segment(page.extract_text() or "", {"page": number}) for number, page in enumerate(reader.pages, 1)]
    metadata = {"page_count": len(reader.pages)}
    if reader.metadata:
        metadata["pdf_metadata"] = {str(key): str(value) for key, value in reader.metadata.items() if value is not None}
    return Extraction(segments, metadata, f"pypdf/{EXTRACTOR_VERSION}")


def _extract_docx(data: bytes) -> Extraction:
    try:
        from docx import Document
    except ImportError as exc:
        raise OptionalDependencyError("DOCX extraction requires python-docx") from exc
    document = Document(BytesIO(data))
    segments = [Segment(paragraph.text, {"paragraph": number})
                for number, paragraph in enumerate(document.paragraphs, 1) if paragraph.text.strip()]
    metadata = {
        "paragraph_count": len(document.paragraphs),
        "core_properties": {
            key: str(value) for key in ("author", "created", "last_modified_by", "modified", "subject", "title")
            if (value := getattr(document.core_properties, key, None)) is not None
        },
    }
    return Extraction(segments, metadata, f"python-docx/{EXTRACTOR_VERSION}")


@lru_cache(maxsize=2)
def _load_captioner(model_id: str):
    try:
        from transformers import pipeline
    except ImportError as exc:
        raise OptionalDependencyError("image captions require the AI dependencies") from exc
    return pipeline("image-to-text", model=model_id)


def _caption_image(image: Any, model_id: str) -> str:
    captioner = _load_captioner(model_id)
    result = captioner(image, max_new_tokens=48)
    return str(result[0]["generated_text"]).strip()


def _extract_image(data: bytes, caption_images: bool, caption_model: str) -> Extraction:
    try:
        from PIL import ExifTags, Image
    except ImportError as exc:
        raise OptionalDependencyError("image metadata requires Pillow") from exc
    with Image.open(BytesIO(data)) as image:
        exif = image.getexif()
        metadata = {
            "format": image.format, "width": image.width, "height": image.height,
            "mode": image.mode,
            "exif": {ExifTags.TAGS.get(key, str(key)): str(value) for key, value in exif.items()},
        }
        caption = _caption_image(image.convert("RGB"), caption_model) if caption_images else ""
    if caption:
        metadata.update({"caption": caption, "caption_model": caption_model})
    searchable = " ".join(filter(None, (
        caption,
        f"Image {metadata['format'] or ''} {metadata['width']}x{metadata['height']} {metadata['mode']}",
        " ".join(f"{key}: {value}" for key, value in metadata["exif"].items()),
    )))
    return Extraction([Segment(searchable, {"type": "image"})] if searchable else [], metadata,
                      f"pillow/{EXTRACTOR_VERSION}" + (f"+caption:{caption_model}" if caption else ""))


def extract_entity(
    entity: dict[str, Any], *, caption_images: bool = False,
    caption_model: str = DEFAULT_CAPTION_MODEL,
) -> Extraction:
    if entity["kind"] in {"note", "upsert"}:
        text = "\n\n".join(value for value in (entity.get("title", ""), entity.get("text", "")) if value)
        return Extraction([Segment(text, {"type": "note"})], {"tags": entity.get("tags", [])}, f"note/{EXTRACTOR_VERSION}")

    blob = entity.get("blob")
    if not blob:
        raise ValueError("document has no blob metadata")
    path = db.blob_path(blob["sha256"])
    if not path.is_file():
        raise FileNotFoundError(f"blob {blob['sha256']} is not available locally")
    data = path.read_bytes()
    mime = (entity.get("mime_type") or blob.get("mime_type") or "").lower()
    filename = (entity.get("filename") or blob.get("original_name") or "").lower()
    if mime == "application/pdf" or filename.endswith(".pdf"):
        return _extract_pdf(data)
    if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or filename.endswith(".docx"):
        return _extract_docx(data)
    if mime.startswith("image/") or filename.endswith((".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff")):
        return _extract_image(data, caption_images, caption_model)
    if mime.startswith("text/") or filename.endswith((".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml")):
        return Extraction([Segment(_decode_text(data), {"type": "document"})], {}, f"text/{EXTRACTOR_VERSION}")
    raise ValueError(f"unsupported document type: {mime or filename or 'unknown'}")


def chunk_segments(
    segments: Iterable[Segment], source_event_id: str,
    max_chars: int = CHUNK_CHARS, overlap_chars: int = CHUNK_OVERLAP_CHARS,
) -> list[dict[str, Any]]:
    if max_chars <= 0 or overlap_chars < 0 or overlap_chars >= max_chars:
        raise ValueError("chunk sizes must satisfy 0 <= overlap < max_chars")
    chunks: list[dict[str, Any]] = []
    ordinal = 0
    absolute_offset = 0
    for segment in segments:
        text = re.sub(r"[ \t]+", " ", segment.text).strip()
        if not text:
            continue
        start = 0
        while start < len(text):
            target = min(start + max_chars, len(text))
            end = target
            if target < len(text):
                boundary = max(text.rfind("\n", start + max_chars // 2, target),
                               text.rfind(". ", start + max_chars // 2, target))
                if boundary > start:
                    end = boundary + 1
            chunk_text = text[start:end].strip()
            if chunk_text:
                chunk_id = hashlib.sha256(
                    f"{source_event_id}:{CHUNKER_VERSION}:{ordinal}:{chunk_text}".encode("utf-8")
                ).hexdigest()
                chunks.append({
                    "chunk_id": chunk_id, "ordinal": ordinal, "text": chunk_text,
                    "start_char": absolute_offset + start, "end_char": absolute_offset + end,
                    "locator": segment.locator,
                })
                ordinal += 1
            if end >= len(text):
                break
            start = max(start + 1, end - overlap_chars)
        absolute_offset += len(text) + 2
    return chunks


def _rebuild_index(
    *, embed: bool = False, caption_images: bool = False,
    embedding_model: str = DEFAULT_EMBEDDING_MODEL,
    caption_model: str = DEFAULT_CAPTION_MODEL,
    force: bool = False,
) -> dict[str, Any]:
    entities = db.current_entities()
    versions = {entity["entity_id"]: entity["_version"] for entity in entities}
    removed = db.remove_stale_derived(versions)
    indexed = skipped = failed = 0
    failures: list[dict[str, str]] = []
    now = datetime.now(timezone.utc).isoformat()
    for entity in entities:
        if not force and db.derived_version(entity["entity_id"]) == entity["_version"]:
            skipped += 1
            continue
        try:
            extraction = extract_entity(entity, caption_images=caption_images, caption_model=caption_model)
            identity_text = " ".join(filter(None, (
                entity.get("title"), entity.get("filename"),
                " ".join(entity.get("tags", [])),
            )))
            if identity_text:
                extraction.segments.insert(0, Segment(identity_text, {"type": "metadata"}))
            chunks = chunk_segments(extraction.segments, entity["_version"])
            metadata = {**extraction.metadata, "chunker_id": CHUNKER_VERSION,
                        "title": entity.get("title"), "filename": entity.get("filename")}
            db.replace_derived_document(entity["entity_id"], entity["_version"], extraction.extractor_id,
                                        extraction.text, metadata, chunks, now)
            indexed += 1
        except Exception as exc:
            db.replace_derived_document(entity["entity_id"], entity["_version"], EXTRACTOR_VERSION,
                                        "", {}, [], now, status="error", error=str(exc))
            failures.append({"entity_id": entity["entity_id"], "error": str(exc)})
            failed += 1
    embedded = build_embeddings(embedding_model) if embed else 0
    return {"indexed": indexed, "skipped": skipped, "failed": failed, "removed": removed,
            "embedded_chunks": embedded, "failures": failures, "status": db.index_status()}


def rebuild_index(
    *, embed: bool = False, caption_images: bool = False,
    embedding_model: str = DEFAULT_EMBEDDING_MODEL,
    caption_model: str = DEFAULT_CAPTION_MODEL,
    force: bool = False,
) -> dict[str, Any]:
    with REBUILD_LOCK:
        return _rebuild_index(
            embed=embed, caption_images=caption_images, embedding_model=embedding_model,
            caption_model=caption_model, force=force,
        )


@lru_cache(maxsize=2)
def _load_embedding_model(model_id: str):
    try:
        from sentence_transformers import SentenceTransformer
    except ImportError as exc:
        raise OptionalDependencyError("embeddings require: pip install -r requirements-ai.txt") from exc
    return SentenceTransformer(model_id)


def build_embeddings(model_id: str = DEFAULT_EMBEDDING_MODEL) -> int:
    chunks = db.indexed_chunks()
    if not chunks:
        db.replace_embeddings(model_id, [])
        return 0
    model = _load_embedding_model(model_id)
    values = model.encode_document([chunk["text"] for chunk in chunks], normalize_embeddings=True)
    packed = []
    for chunk, vector in zip(chunks, values, strict=True):
        floats = array("f", (float(value) for value in vector))
        packed.append((chunk["chunk_id"], len(floats), floats.tobytes()))
    db.replace_embeddings(model_id, packed)
    return len(packed)


def _tokens(text: str) -> list[str]:
    return [token.lower() for token in TOKEN_PATTERN.findall(text)]


def _lexical_scores(query: str, chunks: list[dict[str, Any]]) -> dict[str, float]:
    terms = set(_tokens(query))
    if not terms:
        return {}
    scores = {}
    for chunk in chunks:
        words = _tokens(chunk["text"])
        counts = {term: words.count(term) for term in terms}
        matched = sum(1 for count in counts.values() if count)
        if matched:
            scores[chunk["chunk_id"]] = matched / len(terms) + sum(math.log1p(count) for count in counts.values()) / 10
    return scores


def _semantic_scores(query: str, chunks: list[dict[str, Any]], model_id: str) -> dict[str, float]:
    stored = db.stored_embeddings(model_id)
    if not stored:
        return {}
    model = _load_embedding_model(model_id)
    query_vector = array("f", (float(value) for value in model.encode_query(query, normalize_embeddings=True)))
    scores = {}
    for chunk in chunks:
        record = stored.get(chunk["chunk_id"])
        if record is None:
            continue
        dimensions, data = record
        vector = array("f")
        vector.frombytes(data)
        if dimensions == len(query_vector) == len(vector):
            scores[chunk["chunk_id"]] = sum(left * right for left, right in zip(query_vector, vector))
    return scores


def _rank_fusion(score_sets: list[dict[str, float]]) -> dict[str, float]:
    fused: dict[str, float] = {}
    for scores in score_sets:
        for rank, (chunk_id, _) in enumerate(sorted(scores.items(), key=lambda item: item[1], reverse=True), 1):
            fused[chunk_id] = fused.get(chunk_id, 0.0) + 1.0 / (60 + rank)
    return fused


def search(
    query: str, *, mode: str = "hybrid", limit: int = 10,
    model_id: str = DEFAULT_EMBEDDING_MODEL,
) -> dict[str, Any]:
    chunks = db.indexed_chunks()
    lexical = _lexical_scores(query, chunks) if mode in {"lexical", "hybrid"} else {}
    semantic = _semantic_scores(query, chunks, model_id) if mode in {"semantic", "hybrid"} else {}
    if mode == "semantic" and not semantic:
        raise ValueError("no embeddings are available; rebuild the index with embed=true")
    combined = lexical if mode == "lexical" else semantic if mode == "semantic" else _rank_fusion([lexical, semantic])
    by_id = {chunk["chunk_id"]: chunk for chunk in chunks}
    results = []
    for chunk_id, score in sorted(combined.items(), key=lambda item: item[1], reverse=True)[:limit]:
        chunk = by_id[chunk_id]
        entity = db.current_entity(chunk["entity_id"])
        if entity is None or entity["_version"] != chunk["source_event_id"]:
            continue
        citation = {
            "entity_id": chunk["entity_id"], "source_event_id": chunk["source_event_id"],
            "title": entity.get("title") or entity.get("filename"),
            "filename": entity.get("filename"), "locator": chunk["locator"],
            "start_char": chunk["start_char"], "end_char": chunk["end_char"],
        }
        results.append({"score": score, "text": chunk["text"], "citation": citation})
    return {"query": query, "mode": mode, "model_id": model_id if semantic else None, "results": results}
